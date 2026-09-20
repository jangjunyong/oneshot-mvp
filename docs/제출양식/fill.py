# -*- coding: utf-8 -*-
"""기능설명서 양식(원본) PPTX 에 내용을 채워 작성본을 만든다. (2026-09-18 v2 — 현행 제품 "기획안 팩트체크")

원본은 건드리지 않는다. 표의 셀 텍스트만 갈아 끼우고 서식은 첫 run 의 것을 물려받되, 본문 칸은 글꼴 크기를 명시해
칸을 넘치지 않게 한다(양식 본문 칸의 기본 18pt 상속은 두 줄이면 넘친다). 행 높이는 표 전체 높이 안에서 재배분한다.
숫자는 docs/기능설명서.md(= scripts/spec-numbers.mjs 생성값)와 같은 값만 쓴다.

이미지: 흐름도 4장은 4열(단계마다 한 열), 대표/상세는 슬라이드 9. IMAGES 가 비어 있으면 자리표시자 문자열.
검사: 끝의 check() 가 슬라이드 수·자리표시자·옛 API명·해시태그·미기입 행·URL 을 세어 결과.txt 에 적는다.

실행: python fill.py            → 기능설명서_기획안팩트체크_작성본.pptx
"""
import copy
import io
import os
import re
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN

BLACK = RGBColor(0x1A, 0x1A, 0x1A)

SRC = "기능설명서양식(원본)_(작성용).pptx"
DST = "기능설명서_기획안팩트체크_작성본.pptx"

# ── 사람이 정하는 값 (D0-A) ──────────────────────────────
SERVICE_NAME = "기획안 팩트체크 — 지자체 축제 예산 심의용 실측 검증"
TEAM = "장준용 (숭실대학교 AI소프트웨어학부)"
SITE = "https://oneshot-mvp.vercel.app"
CALL_COUNT_KT = ""      # data.go.kr 마이페이지의 호출건수. 예: "8,921건". 모르면 "" (문장이 빠진다)
CALL_COUNT_FEST = ""
C = "캡처/"
IMAGES = {
    "flow1": [C + "f1_1_home.png", C + "f1_2_underlay.png", C + "f1_3_venue.png", C + "f1_4_panel.png"],
    "flow2": [C + "f2_1_verdict.png", C + "f2_2_stages.png", C + "f2_3_range.png", C + "f2_4_first.png"],
    "flow3": [C + "f3_1_other.png", C + "f3_2_budget.png", C + "f3_3_attrib.png", C + "f3_4_report.png"],
    "flow4": [C + "f4_1_datause.png", C + "f4_2_evidence.png", C + "f4_3_twins.png", C + "f4_4_history.png"],
    "hero": C + "hero_wide.png",
    "detail": C + "detail_4up.png",
}

URL_GUNPO = (SITE + "/check?name=%EA%B5%B0%ED%8F%AC%EC%B2%A0%EC%AD%89%EC%B6%95%EC%A0%9C&sido=%EA%B2%BD%EA%B8%B0&sigungu=%EA%B5%B0%ED%8F%AC%EC%8B%9C"
             "&n=600000&basis=period&counting=personDays&budget=100000&start=2027-04-17&end=2027-04-25"
             "&h1s=2024-04-20&h1e=2024-04-28&h2s=2025-04-19&h2e=2025-04-27&h3s=2026-04-18&h3e=2026-04-26&theme=2&acc=4")
URL_HWACHEON = (SITE + "/check?name=%ED%99%94%EC%B2%9C%EC%82%B0%EC%B2%9C%EC%96%B4%EC%B6%95%EC%A0%9C&sido=%EA%B0%95%EC%9B%90&sigungu=%ED%99%94%EC%B2%9C%EA%B5%B0"
                "&n=1860000&basis=period&counting=personDays&pop=2.3&start=2027-01-09&end=2027-01-31"
                "&h1s=2024-01-06&h1e=2024-01-28&h2s=2025-01-11&h2e=2025-02-02&h3s=2026-01-10&h3e=2026-02-01")

# ── 숫자 (docs/기능설명서.md 와 같은 값) ────────────────
N_FEST = "619건"
N_DAYS = "2,658일"
N_SGG = "272개"
N_ROWS = "697,403행"
PERIOD = "2019-05-01~2026-08-09"
LOO = "정밀도 65.8% · 재현율 61.6% · 무작위 대비 3.26배"
BT = "표본 381건, 두 해 채택 325건 중 167건(51.4%)"

HASHTAGS = ["#축제/행사 기획", "#사전 수요 예측", "#보완 피드백 제공", "#데이터 기반"]
OLD_APIS = ["searchKeyword2", "detailIntro2", "detailCommon2"]


def set_cell(cell, text, pt=None):
    """셀 텍스트 교체. 첫 run 서식을 유지한 채 갈아 끼우고, pt 를 주면 글꼴 크기를 고정한다."""
    tf = cell.text_frame
    lines = text.split("\n")
    p0 = tf.paragraphs[0]
    proto = p0.runs[0] if p0.runs else None
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    if proto is None:
        tf.text = text
    else:
        proto.text = lines[0]
        from pptx.text.text import _Paragraph
        for line in lines[1:]:
            newp = copy.deepcopy(p0._p)
            p0._p.getparent().append(newp)
            para = _Paragraph(newp, tf)
            for r in list(para.runs[1:]):
                r._r.getparent().remove(r._r)
            para.runs[0].text = line
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.color.rgb = BLACK
            if pt:
                run.font.size = Pt(pt)


def link_urls(cell):
    """셀 문단 안의 https:// 주소마다 하이퍼링크를 건다 (pdf 로 내보내면 클릭 가능). 문단 텍스트를 URL 기준으로 쪼개 run 을 다시 만든다."""
    import re as _re
    from pptx.text.text import _Run
    for para in cell.text_frame.paragraphs:
        runs = list(para.runs)
        if not runs:
            continue
        full = "".join(r.text for r in runs)
        if "https://" not in full:
            continue
        proto = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        parts = [x for x in _re.split(r"(https://\S+)", full) if x != ""]
        proto.text = parts[0]
        if parts[0].startswith("https://"):
            proto.hyperlink.address = parts[0]
        last = proto._r
        for part in parts[1:]:
            el = copy.deepcopy(proto._r)
            last.addnext(el)
            last = el
            ru = _Run(el, para)
            ru.text = part
            if part.startswith("https://"):
                ru.hyperlink.address = part
            else:
                # 복제한 run 의 링크는 지운다
                for h in el.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick"):
                    h.getparent().remove(h)


def table(slide, rows=None, cols=None, idx=0):
    ts = [sh for sh in slide.shapes if sh.has_table]
    if rows is not None:
        ts = [sh for sh in ts if len(sh.table.rows) == rows and len(sh.table.columns) == cols]
    return ts[idx].table


def table_shape(slide, rows, cols):
    return [sh for sh in slide.shapes if sh.has_table and len(sh.table.rows) == rows and len(sh.table.columns) == cols][0]


def drop_guides(slide):
    for sh in list(slide.shapes):
        if sh.has_table:
            continue
        if sh.has_text_frame and ("가이드" in sh.text_frame.text or "슬라이드 삭제" in sh.text_frame.text):
            sh._element.getparent().remove(sh._element)


def fill_pairs(t, items, pt, label_col=2):
    """(이름, 설명) 쌍을 2행씩 채우고 남는 행은 지운다."""
    need = len(items) * 2
    while len(t._tbl.tr_lst) > need:
        t._tbl.remove(t._tbl.tr_lst[-1])
    for i, (name, desc) in enumerate(items):
        set_cell(t.cell(i * 2, label_col), name, pt)
        set_cell(t.cell(i * 2 + 1, label_col), desc, pt)


def set_heights(t, heights):
    for r, h in zip(t.rows, heights):
        r.height = Emu(h)


def put_picture(slide, shape, row_idx, col_idx, path):
    """표 셀 위에 그림을 얹는다(셀 안에는 못 넣는다). 셀의 절대 좌표로 맞추고 width 만 준다."""
    t = shape.table
    top = shape.top + sum(t.rows[r].height for r in range(row_idx))
    left = shape.left + sum(t.columns[c].width for c in range(col_idx))
    width = t.columns[col_idx].width
    height = t.rows[row_idx].height
    m = Emu(50000)
    pic = slide.shapes.add_picture(path, left + m, top + m, width=width - 2 * m)
    if pic.height > height - 2 * m:
        ratio = (height - 2 * m) / pic.height
        pic.height = int(pic.height * ratio)
        pic.width = int(pic.width * ratio)
    # 칸 가운데에 놓는다 (세로로 긴 그림이 왼쪽에 붙지 않게)
    pic.left = int(left + (width - pic.width) / 2)
    pic.top = int(top + (height - pic.height) / 2)
    return pic


prs = Presentation(SRC)
S = prs.slides

# ── 1. 표지 ───────────────────────────────────────────────
t = table(S[0])
set_cell(t.cell(0, 1), TEAM)
set_cell(t.cell(1, 1), SERVICE_NAME)

# ── 2. 서비스 소개 (표 높이 5,654,871 안에서 행 재배분) ──
t = table(S[1], 6, 2)
set_heights(t, [480000, 420000, 820000, 1300000, 560000, 2070000])
set_cell(t.cell(0, 1), SERVICE_NAME, 12)
set_cell(t.cell(1, 1), "웹 서비스 (데스크톱 크롬 권장 · 로그인 없음 · 설치 없음)", 12)
set_cell(
    t.cell(2, 1),
    "시·군 문화관광과 축제 담당 공무원과 축제 위탁사 실무자. 예산을 집행하기 전에 기획안의 예상 방문객·예산이 "
    "근거가 있는지를 심의에 증빙해야 하는데, 기댈 것이 작년 발표치와 경험뿐인 사람.",
    12,
)
set_cell(
    t.cell(3, 1),
    "축제 기획서(PDF)를 올리면 예상 방문객·개최 기간·예산을 옮겨 적고, 그 축제가 지난 회차에 실제로 겪은 한국관광공사 KT "
    "시군구 일별 실측(외지인 배수·순증)과 같은 기준으로 재서 통과/주의/과대/과소/상한 초과/근거 없음을 판정합니다. "
    "내년 배수 구간, 1인당 예산 대조, 같은 시기 경쟁 축제 귀속 경고, 결재 첨부용 검증 보고서 A4 2장을 내고, "
    "행사장 도면 위 보행 시뮬레이션으로 쏠림이 어디서 막히는지를 미리 봅니다. 판정·시뮬에 생성형 모델을 쓰지 않습니다.\n"
    "심사용 경로 — 올릴 기획서가 없으면 이 주소가 판정 화면을 바로 엽니다(업로드 불필요):\n" + URL_GUNPO + "\n"
    "시연 안내 " + SITE + "/judge-guide.pdf · 예비 데모 기획서 " + SITE + "/sample-plan.pdf (첫 화면에 올리면 판정이 열립니다)",
    12,
)
link_urls(t.cell(3, 1))
set_cell(t.cell(4, 1), "과제번호 9번 — 축제 흥행 예보 서비스 (#축제/행사 기획 #사전 수요 예측 #보완 피드백 제공 #데이터 기반)", 12)
set_cell(
    t.cell(5, 1),
    "해결과제 원문은 두 문장입니다. ① 축제 수요 예측 실패 및 주관적 경험 의존형 기획으로 인한 예산 낭비 리스크 "
    "② 대규모 관광객 쏠림에 따른 축제 만족도 저하. 이 서비스는 ①을 판정으로, ②를 도면 시뮬레이션으로 정면에서 다룹니다.\n"
    "발표되는 방문객 수는 산출 방식이 공개되지 않고, 내년 기획안은 그 발표치를 옮겨 적는 데서 시작합니다. "
    "군포철쭉축제 2025 발표 최다일 217,502명은 같은 날 KT 가 센 군포시 전체 체류 262,457명의 83%였습니다. "
    "공사 데이터랩에 이동통신 실측이 이미 있는데 기획 단계에서 쓰이지 않습니다. 그 실측을 담당자가 결재에 들고 갈 수 있는 "
    "판정문으로 옮기면 풀리는 문제라고 봤습니다.",
    12,
)

# ── 3. 기획 방향 + 해시태그 목록 ─────────────────────────
t = table(S[2], 2, 2)
set_heights(t, [3300000, 2270000])
set_cell(
    t.cell(0, 1),
    "기획안의 숫자를 실측으로 판정한다 — 방향은 이것 하나입니다.\n"
    "① 자(尺)는 하나: 이 축제 자신의 지난 회차 실측 배수(평소 대비 외지인). 발표치가 아니라 KT 가 센 값입니다.\n"
    "② 세 단계로 잽니다: 모집단 상한(하드 게이트) → 순증분(보조 신호) → 동일 정의 배수(주 근거). 임계값은 화면에 적습니다.\n"
    "③ 흥행 가능성은 예측값이 아니라 세 실측 위치로 답합니다: 자기 이력 안 위치 · 같은 인구 구간 또래 안 위치 · 축제 신호.\n"
    "④ 판정문에는 보완 문장이 따릅니다(무엇을 다시 잡고 무엇의 출처를 밝힐지). \"안전하다\"는 말은 어느 라벨에도 없습니다.\n"
    "⑤ 화면의 명·원 숫자는 전부 출처(API·값·기간·조회일)가 붙은 셀로만 나가고, 테스트가 매번 그것을 셉니다.\n"
    "⑥ 결과는 A4 두 장으로 인쇄됩니다. 결재는 화면이 아니라 종이로 올라가기 때문입니다.",
    12,
)
set_cell(
    t.cell(1, 1),
    "#축제/행사 기획 — 기획서 PDF 입력 → 판정 → 행사장 도면 보행 시뮬레이션(배치 편집·스트레스 테스트)\n"
    "#사전 수요 예측 — 예상 방문객 3단 판정, 내년 배수 구간, 흥행 가능성의 세 실측 위치\n"
    "#보완 피드백 제공 — 라벨별 보완 문장, 1인당 예산 대조, 개최 기간·요일, 귀속 경고, 검증 보고서 A4 2장\n"
    "#데이터 기반 — 공사 KT 일별 실측 " + N_DAYS + "×" + N_SGG + " 전수 적재, 셀마다 출처, 619건 자기검증, −52주 백테스트",
    12,
)

# ── 4. 해시태그 연계 핵심기능 (헤더 + 4행, 표 높이 4,993,748) ──
drop_guides(S[3])
t = table(S[3], 2, 3)
for _ in range(3):
    t._tbl.append(copy.deepcopy(t._tbl.tr_lst[1]))
set_heights(t, [450000, 1135000, 1135000, 1135000, 1135000])
rows = [
    (
        "#축제/행사 기획",
        "기획서 PDF 입력\n행사장 도면 시뮬레이션\n배치 편집·스트레스 테스트",
        "첫 화면에 기획서 PDF 를 올리면 예상 방문객·기간·지역·예산·지난 회차를 옮겨 적어 판정 화면으로 갑니다(모델은 옮겨 적기 1회, "
        "판정에는 안 씀). 시뮬레이션 탭은 같은 기획안을 물고 열리고, 배치도를 밑그림으로 깔아 출입구·부스·통로를 놓은 뒤 보행자를 흘려 "
        "어디가 막히는지 잽니다. 유입 시나리오·출입구 몫은 입력이고 기본값은 출처 없는 가정이라고 적습니다. 스트레스 테스트는 유입 배수를 "
        "0.5씩 올려 행안부 \"위험\"(5명/㎡)이 60초 넘게 유지되는 첫 배수를 냅니다. 말로 물으면 시뮬 수치로만 답합니다.",
    ),
    (
        "#사전 수요 예측",
        "예상 방문객 3단 판정\n내년 배수 구간\n흥행 가능성의 세 실측 위치",
        "① 모집단 상한: 예상 방문객 ÷ 작년 축제 기간 시군구 전체 체류(0.80 이상 상한 초과 · 0.50 이상 과대) ② 순증분: ÷ 작년 축제가 "
        "실제로 끌어온 순증(3.0 이상 주의 신호) ③ 동일 정의 배수: 기획안이 요구하는 평소 대비 배수 ÷ 이 축제 이력의 최대 배수"
        "(0.70/1.00/1.30 → 과소/통과/주의/과대). 연인원이 실인원 상한을 넘으면 손익분기 회전율(군포 발표 217,502명 → 1.04회, 화천 186만 → "
        "2.93회)을 냅니다. 내년 배수 구간은 이력 연도별 배수의 최소~최대(군포 평균 1.1~1.3배·최대일 1.4~1.6배), 이력이 없으면 같은 인구 구간 "
        "또래의 중앙값~상위 5%. 흥행 가능성은 자기 이력 안 위치·또래 안 위치·연도별 축제 신호(뚜렷함/약함)로 보입니다.",
    ),
    (
        "#보완 피드백 제공",
        "라벨별 보완 문장\n1인당 예산 대조\n귀속 경고 · 검증 보고서 A4 2장",
        "판정 라벨마다 담당자가 다음에 할 일을 문장으로 냅니다(\"예상 방문객을 이 축제 이력 배수 안에서 다시 잡고 원래 숫자의 출처와 산정 "
        "방식을 밝힌다\"). 1인당 예산은 기획안 기준과 작년 실측 순증 기준을 나란히 놓아 오차를 보입니다. 개최 기간·요일은 이력 최대일 요일·"
        "주말 포함·길이로 봅니다. 작년 축제 기간 반경 50km 에 다른 축제가 있었으면(공사 searchFestival2 실시간) \"이 배수는 시군구 배수이고 "
        "이 축제 몫만이 아니다\"라고 경고합니다. 검증 보고서 A4 2장(결론·판정표·구간·보완 / 연도별 실측·출처·정의와 임계)은 브라우저 인쇄로 PDF 가 됩니다.",
    ),
    (
        "#데이터 기반",
        "KT 일별 실측 전수 적재\n셀마다 출처\n619건 자기검증 · 백테스트",
        "공사 TourAPI 데이터랩 서비스(locgoRegnVisitrDDList)로 시군구 " + N_SGG + " × " + PERIOD + "(" + N_DAYS + ") 일별 방문자(현지인·외지인·외국인)를 "
        "개발 기간 안에 전수 호출해 " + N_ROWS + "을 적재했습니다. 판정·구간·백테스트의 유일한 실측 원천이고, 데이터랩 축제 목록 " + N_FEST + "의 배수도 "
        "이 자료로 전부 재계산했습니다. 화면의 명·배·원 숫자는 API·값·기간·조회일 네 속성이 붙은 셀로만 나가고 e2e 테스트가 출처 밖의 \"N명\"을 "
        "셉니다. 실측 근거 화면은 연도별 전후 4주 곡선과 표를 보이고, 619건 leave-one-out(" + LOO + ")과 −52주 근사 백테스트(" + BT + ")를 한계와 함께 적습니다.",
    ),
]
for i, (h, f, d) in enumerate(rows, start=1):
    set_cell(t.cell(i, 0), h, 12)
    set_cell(t.cell(i, 1), f, 10)
    set_cell(t.cell(i, 2), d, 9)

# ── 5~8. 기능 흐름도 (해시태그마다 한 장, 4열 = 4단계) ───
flow = [
    ("#축제/행사 기획", "기획서 입력에서 도면 시뮬레이션까지",
     "기획서 PDF 한 장으로 판정 화면이 열리고, 같은 기획안으로 행사장 도면에 들어가 배치를 고치고 보행자를 흘립니다.",
     ["① 첫 화면에 기획서 PDF 를 올리고 [읽어서 판정하기]",
      "② 판정 화면 상단 [시뮬레이션] 탭 → 배치도 밑그림 깔기(축척 두 점 + 거리)",
      "③ 출입구·부스·통로를 놓고 [시뮬레이션 시작] → 도면 위 보행자 흐름과 밀도",
      "④ 오른쪽 패널: 시각·장내 인원·최대 밀도 등급, 병목 상위·대기열·출입구별 입장. [배수 올려 가며 재생]으로 상한 배수"], "flow1"),
    ("#사전 수요 예측", "예상 방문객 3단 판정과 내년 배수 구간",
     "기획안의 숫자를 이 축제 자신의 지난 회차 실측과 같은 자로 세 번 재고, 내년 배수 구간과 흥행 가능성의 위치를 냅니다.",
     ["① 판정 화면 맨 위 라벨(통과/주의/과대/과소/상한 초과/근거 없음)과 한 문단 판정문",
      "② 판정표: 1단계 모집단 상한 → 2단계 순증분 → 3단계 동일 정의 배수(임계값 표기)",
      "③ 내년 배수 구간 카드(평균·최대일) + 또래 구간 + 자기 이력 연도별 배수·축제 신호",
      "④ 이력이 없으면 \"근거 없음\"과 또래 구간만 낸다 — 지어내지 않는다"], "flow2"),
    ("#보완 피드백 제공", "보완 문장에서 검증 보고서까지",
     "라벨마다 다음에 할 일을 문장으로 내고, 예산·기간·경쟁 축제를 대조한 뒤 결재 첨부용 A4 두 장으로 인쇄합니다.",
     ["① 판정문 아래 보완 문장과 \"이 판정이 말하지 않는 것\" 단서",
      "② 1인당 예산 대조(기획안 기준 vs 작년 실측 순증 기준) · 개최 기간·요일 판정",
      "③ 귀속 경고(작년 반경 50km 다른 축제, 공사 API 실시간)",
      "④ [검증 보고서 두 장] → 브라우저 인쇄로 PDF, 결재에 첨부"], "flow3"),
    ("#데이터 기반", "실측 근거와 자기검증",
     "판정에 쓴 모든 값이 어느 API 의 어느 셀에서 왔는지 되짚어지고, 방식의 적중률과 한계를 화면이 먼저 말합니다.",
     ["① 판정표 아래 \"이 판정에 쓴 공사 데이터\" 표(API·어디에·찍힌 값·호출)",
      "② [일별 곡선과 근거 표] → 연도별 전후 4주 외지인 곡선·배수·순증 표",
      "③ 닮은 과거 축제(619건 재계산) 지도·카드와 또래 분포 띠",
      "④ 접힘 \"자기검증\" → leave-one-out 적중률·−52주 백테스트와 한계"], "flow4"),
]
base = S[4]
for i, (h, f, d, steps, key) in enumerate(flow):
    if i == 0:
        sl = base
    else:
        blank = prs.slides.add_slide(base.slide_layout)
        for sh in base.shapes:
            blank.shapes._spTree.append(copy.deepcopy(sh._element))
        sl = blank
    drop_guides(sl)
    t3 = table(sl, 3, 2)
    for r_, lab in enumerate(["해시태그", "연계기능", "기능설명"]):
        set_cell(t3.cell(r_, 0), "%s%d" % (lab, i + 1))
    set_cell(t3.cell(0, 1), h, 12)
    set_cell(t3.cell(1, 1), f, 12)
    set_cell(t3.cell(2, 1), d, 11)
    t4 = table(sl, 3, 4)
    shp = table_shape(sl, 3, 4)
    imgs = IMAGES.get(key) or []
    for c in range(4):
        if c < len(imgs) and os.path.exists(imgs[c]):
            set_cell(t4.cell(1, c), "", 10)
            put_picture(sl, shp, 1, c, imgs[c])
        else:
            set_cell(t4.cell(1, c), "[캡처 %d]" % (c + 1), 10)
        set_cell(t4.cell(2, c), steps[c], 10)

# ── 9. 이미지 ────────────────────────────────────────────
t = table(S[5], 2, 2)
set_heights(t, [2400000, 3170000])
shp = table_shape(S[5], 2, 2)
if IMAGES.get("hero") and os.path.exists(IMAGES["hero"]):
    set_cell(t.cell(0, 1), "", 10)
    put_picture(S[5], shp, 0, 0 + 1, IMAGES["hero"]) if False else put_picture(S[5], shp, 0, 1, IMAGES["hero"])
else:
    set_cell(t.cell(0, 1), "[대표 이미지 — 판정 화면: 예상 방문객 라벨 + 3단 판정표 + 서비스명]", 11)
if IMAGES.get("detail") and os.path.exists(IMAGES["detail"]):
    set_cell(t.cell(1, 1), "", 10)
    put_picture(S[5], shp, 1, 1, IMAGES["detail"])
else:
    set_cell(
        t.cell(1, 1),
        "[상세 이미지 4장 — ① 첫 화면(기획서 PDF 올리기) ② 판정 화면(판정문·3단 표·공사 데이터 표) ③ 실측 근거(연도별 곡선) ④ 시뮬레이션(도면 위 보행자·병목)]\n"
        "심사용 접속 경로 — 서비스 " + SITE + " (로그인 없음)\n"
        "예비 데모 기획서 " + SITE + "/sample-plan.pdf · 기획안 양식 " + SITE + "/plan-form.pdf · 심사위원 시연 안내 " + SITE + "/judge-guide.pdf\n"
        "완성 판정 화면(PDF 없이 바로 들어가는 주소)\n"
        "· 군포철쭉축제 2027 → 주의\n" + URL_GUNPO + "\n"
        "· 화천산천어축제 2027 → 상한 초과\n" + URL_HWACHEON,
        9,
    )
    for para in t.cell(1, 1).text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT

# ── 10. 공사 OpenAPI (최종 사용분만 — 2종) ───────────────
t = table(S[6], 10, 3)
kt_calls = (" 개발 기간 내 호출 " + CALL_COUNT_KT + "(data.go.kr 인증키 기준).") if CALL_COUNT_KT else " 제출 인증키의 호출 이력으로 확인됩니다."
fest_calls = (" 개발 기간 내 호출 " + CALL_COUNT_FEST + ".") if CALL_COUNT_FEST else ""
apis = [
    ("한국관광공사 TourAPI 4.0 관광 데이터랩 서비스 — 지역별 방문자 수 일별 (DataLabService/locgoRegnVisitrDDList, KT 이동통신 기반)",
     "판정 3단계의 분모(작년 축제 기간 전체 체류·순증·평소 대비 외지인 배수), 자기 이력 연도별 배수, 내년 배수 구간, 축제 신호, "
     "데이터랩 축제 목록 " + N_FEST + " 배수 재계산의 유일한 실측 원천. 시군구 " + N_SGG + " × " + PERIOD + "(" + N_DAYS + ")를 "
     "개발 기간 안에 전수 호출해 " + N_ROWS + "을 빌드 전에 적재하고 판정은 그 적재본으로 냅니다." + kt_calls),
    ("한국관광공사 TourAPI 4.0 국문 관광정보 서비스 — 축제 검색 (searchFestival2)",
     "판정 화면에서 실시간 호출. 작년 축제 기간 반경 50km 안에 등록된 다른 축제를 찾아 \"이 배수는 개최 시군구 배수이고 이 축제 몫만이 "
     "아니다\"라는 귀속 경고를 냅니다(군포 2026 기준 등록 행사 36건, 실시간 값). 호출이 실패해도 판정은 그대로 서고 실패 사실을 표시합니다." + fest_calls),
]
fill_pairs(t, apis, 11)

# ── 11. 기타 데이터 (최종 사용분만) ─────────────────────
drop_guides(S[7])
t = table(S[7], 6, 3)
while len(t._tbl.tr_lst) < 10:
    t._tbl.append(copy.deepcopy(t._tbl.tr_lst[2]))
    t._tbl.append(copy.deepcopy(t._tbl.tr_lst[3]))
for i in range(0, 10, 2):
    set_cell(t.cell(i, 0), str(i // 2 + 1))
    set_cell(t.cell(i, 1), "데이터명")
    set_cell(t.cell(i + 1, 1), "상세설명")
etc = [
    ("한국관광 데이터랩 축제 목록 (공사 웹, 정적 수집 · OpenAPI 아님)",
     "전국 축제 " + N_FEST + "의 이름·지역·기간·테마. 배수는 목록 값이 아니라 위 KT 일별 자료로 619/619 재계산. 또래 구간과 닮은 과거 축제(보조 근거)의 표본."),
    ("행정안전부 주민등록 인구 (data.go.kr 15097972, 2026-08 기준)",
     "시군구 인구. 또래(같은 인구 구간) 구간의 기준이고 닮음 판정 다섯 축 중 하나. 담당자 입력 → 행안부 표 → 없음 순."),
    ("OpenStreetMap 부지·도로·산책로 + 군포철쭉축제 2026 공식 안내도",
     "행사장 도면(군포 1건)의 바닥과 부스 배치. 부스 좌표는 georeference 전 가정이라고 화면에 적습니다."),
    ("국토교통부 브이월드 지도 타일 (WMTS) · Natural Earth 1:10m (퍼블릭 도메인)",
     "도면 화면의 위성·일반 배경(대조용)과 닮은 축제 지도의 해안선. 판정에는 쓰지 않습니다."),
    ("우주항공청 월력요항 (공휴일, 2018~2027)",
     "축제 신호 계산에서 설·추석 연휴를 평소 창에서 제외하고 겹치면 표시. 판정 배수에는 넣지 않습니다."),
]
fill_pairs(t, etc, 10)

# ── 12. 차별성 · 발전계획 (표 높이 5,074,306) ────────────
t = table(S[8], 2, 2)
set_heights(t, [2540000, 2530000])
set_cell(
    t.cell(0, 1),
    "첫째, 발표치가 아니라 실측이 판정합니다. 자(尺)는 이 축제 자신이 지난 회차에 겪은 공사 KT 일별 실측 배수이고, 기획안의 숫자를 그 자로 "
    "세 번 잽니다. 예상 방문객 수를 새로 만들어 내지 않으므로 틀릴 수 있는 예측기가 되지 않습니다.\n"
    "둘째, 공사 데이터를 한 번 부르는 것이 아니라 전수로 적재했습니다(시군구 " + N_SGG + " × 최장 " + N_DAYS + ", 시군구마다 시작일이 달라 합계 " + N_ROWS + "). 데이터랩 축제 목록 " + N_FEST + "의 "
    "배수도 같은 자료로 전부 다시 계산해, 또래 비교와 자기검증(" + LOO + ")이 같은 정의 위에 섭니다.\n"
    "셋째, 출력이 결재 문서입니다. 판정문 + 보완 문장 + 검증 보고서 A4 2장. 담당자가 예산 심의에서 \"왜 이 숫자인가\"에 댈 근거입니다.\n"
    "넷째, 화면이 자기 한계를 먼저 말합니다. 명·원 숫자는 출처 셀로만 나가고(e2e 가 셉니다), 임계값은 정한 값이라고 적으며, 내년 구간의 "
    "적중률은 −52주 근사(" + BT + ")로만 쟀다고 카드가 말합니다. 판정·시뮬에 생성형 모델을 쓰지 않아 같은 입력에 늘 같은 답입니다.\n"
    "완성 판정 화면(클릭) — 군포철쭉축제 2027 → 주의: " + URL_GUNPO + "\n"
    "화천산천어축제 2027 → 상한 초과: " + URL_HWACHEON,
    10,
)
link_urls(t.cell(0, 1))
for para in t.cell(0, 1).text_frame.paragraphs:
    if "https://" in "".join(r.text for r in para.runs):
        for r in para.runs:
            r.font.size = Pt(7.5)
set_cell(
    t.cell(1, 1),
    "1) 공사 소비·내비게이션 검색·SNS 지수(AreaTarResDemService) 연동 — 활용신청 진행 중. 예산 항목별(먹거리·주차·홍보) 근거를 같은 판정문 안에 늘립니다.\n"
    "2) 다년 기간 수집 — 문화체육관광부 지역축제 파일의 회차별 기간을 모아 −52주 근사가 아니라 실제 다년 표본으로 내년 구간의 적중률을 다시 재고, "
    "임계값도 군포 한 건이 아니라 다도시로 다시 맞춥니다.\n"
    "3) 사후 평가 루프 — 개최 뒤 실측이 데이터랩에 쌓이면 기획안 대비 실적을 같은 자로 다시 재고 다음 기획안 판정에 넘깁니다. 담당자 업무가 "
    "결과보고서 → 내년 기획안으로 이어지는 흐름 그대로입니다.\n"
    "4) 안전관리계획 서식 연동 — 행정안전부 「지역축제장 안전관리 매뉴얼」의 \"최대 수용인원\"·\"관람객 동선\" 칸에 판정표와 도면 시뮬레이션 결과를 "
    "그대로 붙여, 매뉴얼에 없는 통로 폭·밀집도 산출 근거의 빈칸을 채웁니다.\n"
    "5) 배치도 자동 인식 — 지금은 배치도를 밑그림으로 깔고 따라 그립니다. 실물 배치도로 정확도를 잰 뒤 모델 초안 생성을 붙입니다.\n"
    "운영: 정적 사이트 + 서버리스(Vercel)라 고정비가 없고, 공사 데이터는 월 1회 프리페치(개발계정 1,000회/일 안)로 갱신합니다. 운영계정 전환 시 실시간 호출 한도 문제도 사라집니다.",
    10,
)

# 복제한 흐름도 3장은 덱 끝에 붙으므로 5번 슬라이드 뒤(6·7·8)로 옮긴다
_ids = prs.slides._sldIdLst
for k, el in enumerate(list(_ids)[-3:]):
    _ids.remove(el)
    _ids.insert(5 + k, el)

prs.save(DST)


# ── 검사 ─────────────────────────────────────────────────
def check(path):
    p = Presentation(path)
    texts = []
    for s in p.slides:
        for sh in s.shapes:
            if sh.has_table:
                for r in range(len(sh.table.rows)):
                    for c in range(len(sh.table.columns)):
                        texts.append(sh.table.cell(r, c).text)
            elif sh.has_text_frame:
                texts.append(sh.text_frame.text)
    all_text = "\n".join(texts)
    out = []
    out.append("slides=%d (expect 12)" % len(p.slides))
    out.append("placeholders=%d" % len(re.findall(r"\[(캡처 \d|대표 이미지|상세 이미지)", all_text)))
    out.append("old_api_hits=%d (expect 0)" % sum(all_text.count(a) for a in OLD_APIS))
    out.append("hashtags_ok=%s" % all(h in all_text for h in HASHTAGS))
    out.append("guide_boxes=%d (expect 0)" % sum(1 for t_ in texts if "가이드" in t_ and "예시" in t_))
    for si, s in enumerate(p.slides):
        tbs = [sh.table for sh in s.shapes if sh.has_table]
        if not tbs or len(tbs[0].columns) != 3 or tbs[0].cell(0, 1).text.strip() not in ("API명", "데이터명"):
            continue
        tb = tbs[0]
        empty = sum(1 for r in range(len(tb.rows)) if tb.cell(r, 2).text.strip() == "")
        out.append("slide%d(%s)_empty_rows=%d (expect 0)" % (si + 1, tb.cell(0, 1).text.strip(), empty))
    out.append("urls=%d (expect >=2)" % all_text.count(SITE + "/check?"))
    out.append("service_name_ok=%s" % (SERVICE_NAME in all_text))
    return "\n".join(out)


report = "saved %s / slides %d\n%s\n" % (DST, len(prs.slides), check(DST))
io.open("결과.txt", "w", encoding="utf-8").write(report)
sys.stdout.buffer.write(report.encode("utf-8"))
